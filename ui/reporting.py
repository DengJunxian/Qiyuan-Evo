"""Helpers for generating downloadable policy and replay report artifacts."""

from __future__ import annotations

import hashlib
import json
import re
from datetime import datetime
from io import BytesIO
from pathlib import Path
from typing import Any, Dict, List, Mapping, Optional, Tuple

import pandas as pd

try:
    from docx import Document
    from docx.enum.section import WD_SECTION
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Mm, Pt, RGBColor

    DOCX_EXPORT_AVAILABLE = True
except Exception:
    Document = Any  # type: ignore[assignment]
    WD_SECTION = None  # type: ignore[assignment]
    WD_ALIGN_PARAGRAPH = None  # type: ignore[assignment]
    OxmlElement = None  # type: ignore[assignment]
    qn = None  # type: ignore[assignment]
    Mm = None  # type: ignore[assignment]
    Pt = None  # type: ignore[assignment]
    RGBColor = None  # type: ignore[assignment]
    DOCX_EXPORT_AVAILABLE = False

try:
    from reportlab.lib import colors
    from reportlab.lib.enums import TA_CENTER, TA_JUSTIFY, TA_LEFT
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
    from reportlab.lib.units import mm
    from reportlab.pdfbase.cidfonts import UnicodeCIDFont
    from reportlab.pdfbase.pdfmetrics import registerFont
    from reportlab.platypus import PageBreak, Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

    PDF_EXPORT_AVAILABLE = True
except Exception:
    colors = None  # type: ignore[assignment]
    TA_CENTER = TA_JUSTIFY = TA_LEFT = None  # type: ignore[assignment]
    A4 = None  # type: ignore[assignment]
    ParagraphStyle = None  # type: ignore[assignment]
    getSampleStyleSheet = None  # type: ignore[assignment]
    mm = None  # type: ignore[assignment]
    UnicodeCIDFont = None  # type: ignore[assignment]
    registerFont = None  # type: ignore[assignment]
    PageBreak = Paragraph = SimpleDocTemplate = Spacer = Table = TableStyle = None  # type: ignore[assignment]
    PDF_EXPORT_AVAILABLE = False


def safe_slug(text: str, fallback: str = "report", max_length: int = 36) -> str:
    normalized = re.sub(r"\s+", "-", str(text or "").strip().lower())
    normalized = re.sub(r"[^a-z0-9\-_]+", "-", normalized).strip("-_")
    return (normalized or fallback)[:max_length]


def serializable_payload(payload: Dict[str, Any]) -> Dict[str, Any]:
    return json.loads(json.dumps(payload, ensure_ascii=False, default=str))


def export_capabilities() -> Dict[str, Any]:
    missing: List[str] = []
    if not DOCX_EXPORT_AVAILABLE:
        missing.append("python-docx")
    if not PDF_EXPORT_AVAILABLE:
        missing.append("reportlab")
    return {
        "docx": DOCX_EXPORT_AVAILABLE,
        "pdf": PDF_EXPORT_AVAILABLE,
        "missing_dependencies": missing,
    }


def official_report_meta(report_type: str, title: str) -> Dict[str, str]:
    now = datetime.now()
    stamp = now.strftime("%Y%m%d-%H%M%S")
    prefix = {
        "policy_lab": "CIVITAS-POL",
        "history_replay": "CIVITAS-HIS",
        "realism_report": "CIVITAS-REAL",
    }.get(report_type, "CIVITAS-RPT")
    recipient = {
        "policy_lab": "政策研究会商组、相关业务部门",
        "history_replay": "政策评估组、历史复盘与校准团队",
    }.get(report_type, "内部研究与评估团队")
    return {
        "report_no": f"{prefix}-{stamp}",
        "date_cn": now.strftime("%Y年%m月%d日"),
        "generated_at": now.isoformat(timespec="seconds"),
        "issuer": "Civitas 政策仿真系统",
        "classification": "内部研判材料",
        "recipient": recipient,
        "title": title,
    }


def _parse_markdown(markdown_text: str) -> List[Tuple[str, str]]:
    blocks: List[Tuple[str, str]] = []
    for raw_line in markdown_text.splitlines():
        line = raw_line.rstrip()
        if not line:
            blocks.append(("blank", ""))
        elif line.startswith("# "):
            blocks.append(("h1", line[2:].strip()))
        elif line.startswith("## "):
            blocks.append(("h2", line[3:].strip()))
        elif line.startswith("### "):
            blocks.append(("h3", line[4:].strip()))
        elif line.startswith("- "):
            blocks.append(("bullet", line[2:].strip()))
        else:
            blocks.append(("text", line.strip()))
    return blocks


def _content_blocks(markdown_text: str, report_meta: Dict[str, str]) -> List[Tuple[str, str]]:
    blocks = _parse_markdown(markdown_text)
    if blocks and blocks[0][0] == "h1" and blocks[0][1] == report_meta["title"]:
        return blocks[1:]
    return blocks


def _set_run_font(run, size_pt: float, *, bold: bool = False, color: str | None = None) -> None:
    run.font.size = Pt(size_pt)
    run.bold = bold
    run.font.name = "宋体"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
    if color:
        run.font.color.rgb = RGBColor.from_string(color)


def _append_page_number(paragraph) -> None:
    run = paragraph.add_run()
    fld_begin = OxmlElement("w:fldChar")
    fld_begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = "PAGE"
    fld_end = OxmlElement("w:fldChar")
    fld_end.set(qn("w:fldCharType"), "end")
    run._r.append(fld_begin)
    run._r.append(instr)
    run._r.append(fld_end)


def _build_docx_cover(document: Document, report_meta: Dict[str, str]) -> None:
    title = document.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.paragraph_format.space_before = Pt(120)
    title_run = title.add_run(report_meta["title"])
    _set_run_font(title_run, 22, bold=True, color="17304F")

    subtitle = document.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle_run = subtitle.add_run("正式汇报材料")
    _set_run_font(subtitle_run, 14, bold=True, color="305178")

    meta_table = document.add_table(rows=5, cols=2)
    meta_table.style = "Table Grid"
    rows = [
        ("报告编号", report_meta["report_no"]),
        ("报送对象", report_meta["recipient"]),
        ("材料性质", report_meta["classification"]),
        ("生成日期", report_meta["date_cn"]),
        ("出具单位", report_meta["issuer"]),
    ]
    for idx, (label, value) in enumerate(rows):
        cells = meta_table.rows[idx].cells
        cells[0].text = label
        cells[1].text = value
        for cell in cells:
            for paragraph in cell.paragraphs:
                for run in paragraph.runs:
                    _set_run_font(run, 11)

    note = document.add_paragraph()
    note.alignment = WD_ALIGN_PARAGRAPH.CENTER
    note.paragraph_format.space_before = Pt(24)
    note_run = note.add_run("本材料用于政策研判、会商沟通和正式汇报底稿。")
    _set_run_font(note_run, 11, color="4B6587")


def _build_docx_body(document: Document, markdown_text: str, report_meta: Dict[str, str]) -> None:
    section = document.sections[0]
    section.page_width = Mm(210)
    section.page_height = Mm(297)
    section.top_margin = Mm(25)
    section.bottom_margin = Mm(22)
    section.left_margin = Mm(24)
    section.right_margin = Mm(20)
    section.different_first_page_header_footer = True

    _build_docx_cover(document, report_meta)

    content_section = document.add_section(WD_SECTION.NEW_PAGE)
    content_section.page_width = Mm(210)
    content_section.page_height = Mm(297)
    content_section.top_margin = Mm(22)
    content_section.bottom_margin = Mm(18)
    content_section.left_margin = Mm(24)
    content_section.right_margin = Mm(20)

    header = content_section.header.paragraphs[0]
    header.alignment = WD_ALIGN_PARAGRAPH.LEFT
    header_run = header.add_run(f"{report_meta['issuer']}    {report_meta['report_no']}")
    _set_run_font(header_run, 9, color="4B6587")

    footer = content_section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    footer_run = footer.add_run(f"{report_meta['classification']}  |  第 ")
    _set_run_font(footer_run, 9, color="4B6587")
    _append_page_number(footer)
    footer_tail = footer.add_run(" 页")
    _set_run_font(footer_tail, 9, color="4B6587")

    for kind, text in _content_blocks(markdown_text, report_meta):
        if kind == "blank":
            document.add_paragraph()
            continue

        paragraph = document.add_paragraph()
        if kind == "h1":
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            run = paragraph.add_run(text)
            _set_run_font(run, 18, bold=True, color="17304F")
        elif kind == "h2":
            paragraph.paragraph_format.space_before = Pt(10)
            run = paragraph.add_run(text)
            _set_run_font(run, 14, bold=True, color="17304F")
        elif kind == "h3":
            run = paragraph.add_run(text)
            _set_run_font(run, 12, bold=True, color="305178")
        elif kind == "bullet":
            paragraph.style = document.styles["List Bullet"]
            run = paragraph.add_run(text)
            _set_run_font(run, 10.5)
        else:
            paragraph.paragraph_format.first_line_indent = Pt(21)
            paragraph.paragraph_format.line_spacing = 1.5
            paragraph.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            run = paragraph.add_run(text)
            _set_run_font(run, 10.5)


def _build_docx_bytes(markdown_text: str, report_meta: Dict[str, str]) -> bytes:
    if not DOCX_EXPORT_AVAILABLE:
        raise RuntimeError("python-docx is not available")
    document = Document()
    _build_docx_body(document, markdown_text, report_meta)
    output = BytesIO()
    document.save(output)
    return output.getvalue()


def _pdf_styles():
    if not PDF_EXPORT_AVAILABLE:
        raise RuntimeError("reportlab is not available")
    registerFont(UnicodeCIDFont("STSong-Light"))
    styles = getSampleStyleSheet()
    return {
        "cover_title": ParagraphStyle(
            "CoverTitle",
            parent=styles["Title"],
            fontName="STSong-Light",
            fontSize=22,
            leading=28,
            alignment=TA_CENTER,
            textColor=colors.HexColor("#17304F"),
        ),
        "cover_subtitle": ParagraphStyle(
            "CoverSubtitle",
            parent=styles["Normal"],
            fontName="STSong-Light",
            fontSize=13,
            leading=18,
            alignment=TA_CENTER,
            textColor=colors.HexColor("#305178"),
        ),
        "h1": ParagraphStyle(
            "ReportH1",
            parent=styles["Heading1"],
            fontName="STSong-Light",
            fontSize=17,
            leading=24,
            alignment=TA_CENTER,
            textColor=colors.HexColor("#17304F"),
        ),
        "h2": ParagraphStyle(
            "ReportH2",
            parent=styles["Heading2"],
            fontName="STSong-Light",
            fontSize=13,
            leading=19,
            spaceBefore=10,
            spaceAfter=4,
            alignment=TA_LEFT,
            textColor=colors.HexColor("#17304F"),
        ),
        "h3": ParagraphStyle(
            "ReportH3",
            parent=styles["Heading3"],
            fontName="STSong-Light",
            fontSize=11.5,
            leading=16,
            spaceBefore=6,
            spaceAfter=2,
            alignment=TA_LEFT,
            textColor=colors.HexColor("#305178"),
        ),
        "body": ParagraphStyle(
            "ReportBody",
            parent=styles["BodyText"],
            fontName="STSong-Light",
            fontSize=10.5,
            leading=16,
            alignment=TA_JUSTIFY,
            firstLineIndent=21,
            textColor=colors.HexColor("#1F2B44"),
        ),
        "bullet": ParagraphStyle(
            "ReportBullet",
            parent=styles["BodyText"],
            fontName="STSong-Light",
            fontSize=10.5,
            leading=16,
            leftIndent=14,
            firstLineIndent=-10,
            bulletIndent=0,
            textColor=colors.HexColor("#1F2B44"),
        ),
        "meta": ParagraphStyle(
            "ReportMeta",
            parent=styles["BodyText"],
            fontName="STSong-Light",
            fontSize=10.5,
            leading=15,
            alignment=TA_LEFT,
            textColor=colors.HexColor("#1F2B44"),
        ),
    }


def _draw_pdf_header_footer(canvas, doc, report_meta: Dict[str, str]) -> None:
    if doc.page == 1:
        return
    canvas.saveState()
    canvas.setFont("STSong-Light", 9)
    canvas.setFillColor(colors.HexColor("#4B6587"))
    canvas.drawString(doc.leftMargin, A4[1] - 14 * mm, report_meta["issuer"])
    canvas.drawRightString(A4[0] - doc.rightMargin, A4[1] - 14 * mm, report_meta["report_no"])
    canvas.drawString(doc.leftMargin, 10 * mm, report_meta["classification"])
    canvas.drawCentredString(A4[0] / 2, 10 * mm, f"第 {canvas.getPageNumber() - 1} 页")
    canvas.restoreState()


def _build_pdf_bytes(markdown_text: str, report_meta: Dict[str, str]) -> bytes:
    if not PDF_EXPORT_AVAILABLE:
        raise RuntimeError("reportlab is not available")
    styles = _pdf_styles()
    buffer = BytesIO()
    doc = SimpleDocTemplate(
        buffer,
        pagesize=A4,
        leftMargin=24 * mm,
        rightMargin=20 * mm,
        topMargin=22 * mm,
        bottomMargin=18 * mm,
    )

    story: List[Any] = []
    story.append(Spacer(1, 50 * mm))
    story.append(Paragraph(report_meta["title"], styles["cover_title"]))
    story.append(Spacer(1, 8 * mm))
    story.append(Paragraph("正式汇报材料", styles["cover_subtitle"]))
    story.append(Spacer(1, 18 * mm))

    cover_table = Table(
        [
            ["报告编号", report_meta["report_no"]],
            ["报送对象", report_meta["recipient"]],
            ["材料性质", report_meta["classification"]],
            ["生成日期", report_meta["date_cn"]],
            ["出具单位", report_meta["issuer"]],
        ],
        colWidths=[32 * mm, 118 * mm],
    )
    cover_table.setStyle(
        TableStyle(
            [
                ("FONTNAME", (0, 0), (-1, -1), "STSong-Light"),
                ("FONTSIZE", (0, 0), (-1, -1), 10.5),
                ("TEXTCOLOR", (0, 0), (-1, -1), colors.HexColor("#1F2B44")),
                ("BACKGROUND", (0, 0), (0, -1), colors.HexColor("#EDF2F9")),
                ("GRID", (0, 0), (-1, -1), 0.4, colors.HexColor("#C9D8EF")),
                ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
                ("TOPPADDING", (0, 0), (-1, -1), 6),
                ("BOTTOMPADDING", (0, 0), (-1, -1), 6),
            ]
        )
    )
    story.append(cover_table)
    story.append(Spacer(1, 16 * mm))
    story.append(Paragraph("本材料用于政策研判、会商沟通和正式汇报底稿。", styles["cover_subtitle"]))
    story.append(PageBreak())

    for kind, text in _content_blocks(markdown_text, report_meta):
        if kind == "blank":
            story.append(Spacer(1, 3 * mm))
        elif kind == "h1":
            story.append(Paragraph(text, styles["h1"]))
        elif kind == "h2":
            story.append(Paragraph(text, styles["h2"]))
        elif kind == "h3":
            story.append(Paragraph(text, styles["h3"]))
        elif kind == "bullet":
            story.append(Paragraph(f"• {text}", styles["bullet"]))
        else:
            story.append(Paragraph(text, styles["body"]))

    doc.build(
        story,
        onFirstPage=lambda canvas, doc: None,
        onLaterPages=lambda canvas, doc: _draw_pdf_header_footer(canvas, doc, report_meta),
    )
    return buffer.getvalue()


def write_report_artifacts(
    *,
    root_dir: Path,
    report_type: str,
    title: str,
    markdown_text: str,
    payload: Dict[str, Any],
) -> Dict[str, Any]:
    timestamp = datetime.now().strftime("%Y%m%d-%H%M%S")
    root_dir.mkdir(parents=True, exist_ok=True)
    stem = f"{report_type}_{safe_slug(title, fallback=report_type)}_{timestamp}"
    markdown_path = root_dir / f"{stem}.md"
    json_path = root_dir / f"{stem}.json"
    docx_path = root_dir / f"{stem}.docx"
    pdf_path = root_dir / f"{stem}.pdf"

    clean_payload = serializable_payload(payload)
    report_meta = clean_payload.get("report_meta") or official_report_meta(report_type, title)
    experiment_id = (
        clean_payload.get("experiment_id")
        or clean_payload.get("report_experiment_id")
        or dict(clean_payload.get("reproducibility") or {}).get("experiment_id")
    )
    if not experiment_id:
        experiment_id = f"exp_{stable_payload_hash({'report_type': report_type, 'title': title, 'seed': dict(clean_payload.get('reproducibility') or {}).get('seed', ''), 'config_hash': clean_payload.get('config_hash', '')})[:20]}"
    clean_payload["experiment_id"] = str(experiment_id)
    clean_payload["report_experiment_id"] = str(experiment_id)
    report_meta["experiment_id"] = str(experiment_id)
    clean_payload["report_meta"] = report_meta

    markdown_path.write_text(markdown_text, encoding="utf-8")
    json_text = json.dumps(clean_payload, ensure_ascii=False, indent=2)
    json_path.write_text(json_text, encoding="utf-8")

    capabilities = export_capabilities()
    disabled_reasons: Dict[str, str] = {}
    docx_bytes: Optional[bytes] = None
    pdf_bytes: Optional[bytes] = None

    if capabilities["docx"]:
        docx_bytes = _build_docx_bytes(markdown_text, report_meta)
        docx_path.write_bytes(docx_bytes)
    else:
        disabled_reasons["docx"] = "python-docx is not installed"

    if capabilities["pdf"]:
        pdf_bytes = _build_pdf_bytes(markdown_text, report_meta)
        pdf_path.write_bytes(pdf_bytes)
    else:
        disabled_reasons["pdf"] = "reportlab is not installed"

    return {
        "stem": stem,
        "timestamp": timestamp,
        "report_meta": report_meta,
        "markdown_path": markdown_path,
        "json_path": json_path,
        "docx_path": docx_path if docx_bytes is not None else None,
        "pdf_path": pdf_path if pdf_bytes is not None else None,
        "markdown_text": markdown_text,
        "json_text": json_text,
        "docx_bytes": docx_bytes,
        "pdf_bytes": pdf_bytes,
        "export_capabilities": capabilities,
        "disabled_reasons": disabled_reasons,
    }


def dataframe_export_bundle(frame: pd.DataFrame, *, stem: str) -> Dict[str, Any]:
    """Build CSV and optional Parquet bytes for report-side data exports."""

    safe_stem = safe_slug(stem, fallback="research_data")
    data = frame.copy() if isinstance(frame, pd.DataFrame) else pd.DataFrame()
    csv_text = data.to_csv(index=False)
    parquet_bytes: Optional[bytes] = None
    disabled_reasons: Dict[str, str] = {}
    try:
        output = BytesIO()
        data.to_parquet(output, index=False)
        parquet_bytes = output.getvalue()
    except Exception as exc:
        disabled_reasons["parquet"] = f"parquet export unavailable: {exc}"
    return {
        "stem": safe_stem,
        "csv_text": csv_text,
        "csv_bytes": csv_text.encode("utf-8-sig"),
        "parquet_bytes": parquet_bytes,
        "disabled_reasons": disabled_reasons,
    }


def _stable_json_text(payload: Mapping[str, Any]) -> str:
    return json.dumps(serializable_payload(dict(payload)), ensure_ascii=False, sort_keys=True, separators=(",", ":"))


def stable_payload_hash(payload: Mapping[str, Any]) -> str:
    return hashlib.sha256(_stable_json_text(payload).encode("utf-8")).hexdigest()


def _format_value(value: Any) -> str:
    if isinstance(value, (dict, list, tuple)):
        return json.dumps(value, ensure_ascii=False, sort_keys=True)
    if isinstance(value, float):
        return f"{value:.6f}"
    return str(value)


def _markdown_table(section: Mapping[str, Any], *, skip: Tuple[str, ...] = ("enabled",)) -> str:
    rows = []
    for key, value in section.items():
        if key in skip:
            continue
        rows.append((key, _format_value(value)))
    if not rows:
        return "_暂无可用指标。_"
    lines = ["| 指标 | 数值 |", "| --- | --- |"]
    lines.extend([f"| {key} | {value} |" for key, value in rows])
    return "\n".join(lines)


def render_realism_report_markdown(payload: Mapping[str, Any]) -> str:
    title = str(payload.get("title", "典型事实拟真评估报告"))
    meta = payload.get("report_meta") or official_report_meta("realism_report", title)
    path_fit = payload.get("path_fit", {})
    micro = payload.get("microstructure_fit", {})
    behavior = payload.get("behavioral_fit", {})
    scorecard = payload.get("replay_scorecard", payload.get("scorecard", {}))
    repro = payload.get("reproducibility", {})
    snapshot = payload.get("snapshot_info", {})
    charts = payload.get("charts", [])

    lines = [
        f"# {meta['title']}",
        "",
        f"- 报告编号：{meta['report_no']}",
        f"- 生成时间：{meta['generated_at']}",
        f"- 随机种子：{repro.get('seed', payload.get('seed', 0))}",
        f"- 配置哈希：{repro.get('config_hash', payload.get('config_hash', ''))}",
        f"- 功能开关：{payload.get('feature_flag', False)}",
        "",
        "## 快照信息",
        "```json",
        json.dumps(snapshot, ensure_ascii=False, indent=2, sort_keys=True),
        "```",
        "",
        "## Path Fit（路径拟合）",
        _markdown_table(path_fit),
        "",
        "## Microstructure Fit（微观结构拟合）",
        _markdown_table(micro),
        "",
        "## Behavioral Fit（行为拟合）",
        _markdown_table(behavior),
        "",
        "## 统一回放评估卡",
        _markdown_table(scorecard.get("path_fit_metrics", {}) if isinstance(scorecard, Mapping) else {}),
        "",
        "## 图表",
    ]
    if charts:
        for chart in charts:
            lines.append(f"- {chart.get('name', 'chart')}: {chart.get('kind', 'unknown')}")
    else:
        lines.append("- 无")
    lines.extend(
        [
            "",
            "## 可复现信息",
            "```json",
            json.dumps(repro, ensure_ascii=False, indent=2, sort_keys=True),
            "```",
        ]
    )
    return "\n".join(lines)


def write_realism_report_artifacts(
    *,
    root_dir: Path,
    title: str,
    payload: Dict[str, Any],
    feature_flag: bool = False,
) -> Dict[str, Any]:
    clean_payload = serializable_payload(payload)
    feature_flags = dict(clean_payload.get("feature_flags") or {})
    feature_flags["stylized_facts_v2"] = bool(feature_flag or clean_payload.get("feature_flag", False))
    clean_payload["feature_flags"] = feature_flags
    clean_payload["feature_flag"] = feature_flags["stylized_facts_v2"]
    clean_payload.setdefault("title", title)
    clean_payload["report_meta"] = clean_payload.get("report_meta") or official_report_meta("realism_report", title)
    clean_payload["config_hash"] = clean_payload.get("config_hash") or stable_payload_hash(
        {
            "title": title,
            "seed": clean_payload.get("reproducibility", {}).get("seed", 0),
            "feature_flag": clean_payload["feature_flag"],
            "snapshot_info": clean_payload.get("snapshot_info", {}),
        }
    )

    markdown_text = render_realism_report_markdown(clean_payload)
    bundle = write_report_artifacts(
        root_dir=root_dir,
        report_type="realism_report",
        title=title,
        markdown_text=markdown_text,
        payload=clean_payload,
    )

    charts_path = bundle["markdown_path"].with_name(f"{bundle['stem']}_charts.json")
    charts_text = json.dumps(clean_payload.get("charts", []), ensure_ascii=False, indent=2, sort_keys=True)
    charts_path.write_text(charts_text, encoding="utf-8")
    bundle["charts_path"] = charts_path
    bundle["charts_text"] = charts_text
    bundle["feature_flags"] = feature_flags
    return bundle


def export_defense_bundle(
    *,
    root_dir: Path,
    bundle_name: str,
    design_chapter_markdown: str,
    realism_payload: Dict[str, Any],
    policy_ab_markdown: str,
    architecture_graph: Dict[str, Any],
    causal_chain_graph: Dict[str, Any],
    defense_outline_markdown: str,
    feature_flags: Optional[Dict[str, Any]] = None,
    compliance_artifacts: Optional[Dict[str, Any]] = None,
    social_propagation_artifacts: Optional[Dict[str, Any]] = None,
    agent_taxonomy_markdown: Optional[str] = None,
    policy_causal_chain: Optional[Dict[str, Any]] = None,
    realism_metrics_csv: Optional[str] = None,
    competition_snapshot: Optional[Dict[str, Any]] = None,
) -> Dict[str, Any]:
    """
    Export one-click defense materials package:
    - design chapter draft
    - realism report
    - policy A/B report
    - architecture / causal chain graph data
    - defense speech outline
    - competition compliance artifacts
    - social propagation export
    """
    root_dir.mkdir(parents=True, exist_ok=True)
    stamp = datetime.now().strftime("%Y%m%d-%H%M%S")
    safe_name = safe_slug(bundle_name, fallback="defense_bundle")
    bundle_root = root_dir / f"{safe_name}_{stamp}"
    bundle_root.mkdir(parents=True, exist_ok=True)

    design_path = bundle_root / "design_chapter_draft.md"
    policy_ab_path = bundle_root / "policy_ab_report.md"
    architecture_path = bundle_root / "architecture_graph.json"
    causal_chain_path = bundle_root / "causal_chain_graph.json"
    outline_path = bundle_root / "defense_outline.md"
    social_path = bundle_root / "social_propagation_report.json"
    agent_taxonomy_path = bundle_root / "agent_taxonomy.md"
    policy_causal_chain_path = bundle_root / "policy_causal_chain.json"
    realism_metrics_path = bundle_root / "realism_metrics.csv"
    competition_snapshot_path = bundle_root / "competition_mode_snapshot.json"
    manifest_path = bundle_root / "bundle_manifest.json"

    design_path.write_text(design_chapter_markdown, encoding="utf-8")
    policy_ab_path.write_text(policy_ab_markdown, encoding="utf-8")
    architecture_path.write_text(json.dumps(architecture_graph, ensure_ascii=False, indent=2, sort_keys=True), encoding="utf-8")
    causal_chain_path.write_text(json.dumps(causal_chain_graph, ensure_ascii=False, indent=2, sort_keys=True), encoding="utf-8")
    outline_path.write_text(defense_outline_markdown, encoding="utf-8")
    if social_propagation_artifacts:
        social_path.write_text(
            json.dumps(serializable_payload(social_propagation_artifacts), ensure_ascii=False, indent=2, sort_keys=True),
            encoding="utf-8",
        )
    if agent_taxonomy_markdown:
        agent_taxonomy_path.write_text(str(agent_taxonomy_markdown), encoding="utf-8")
    if policy_causal_chain is not None:
        policy_causal_chain_path.write_text(
            json.dumps(serializable_payload(policy_causal_chain), ensure_ascii=False, indent=2, sort_keys=True),
            encoding="utf-8",
        )
    if realism_metrics_csv:
        realism_metrics_path.write_text(str(realism_metrics_csv), encoding="utf-8")
    if competition_snapshot is not None:
        competition_snapshot_path.write_text(
            json.dumps(serializable_payload(competition_snapshot), ensure_ascii=False, indent=2, sort_keys=True),
            encoding="utf-8",
        )

    realism_bundle = write_realism_report_artifacts(
        root_dir=bundle_root,
        title=str(realism_payload.get("title", "???????")),
        payload=realism_payload,
        feature_flag=bool((feature_flags or {}).get("stylized_facts_v2", True)),
    )

    manifest = {
        "bundle_name": bundle_name,
        "created_at": datetime.now().isoformat(timespec="seconds"),
        "feature_flags": dict(feature_flags or {}),
        "files": {
            "design_chapter_draft": str(design_path),
            "realism_markdown": str(realism_bundle["markdown_path"]),
            "realism_json": str(realism_bundle["json_path"]),
            "policy_ab_report": str(policy_ab_path),
            "architecture_graph": str(architecture_path),
            "causal_chain_graph": str(causal_chain_path),
            "defense_outline": str(outline_path),
        },
    }
    if realism_bundle.get("docx_path") is not None:
        manifest["files"]["realism_docx"] = str(realism_bundle["docx_path"])
    if realism_bundle.get("pdf_path") is not None:
        manifest["files"]["realism_pdf"] = str(realism_bundle["pdf_path"])
    if realism_bundle.get("disabled_reasons"):
        manifest["export_disabled_reasons"] = serializable_payload(realism_bundle["disabled_reasons"])
    if social_propagation_artifacts:
        manifest["files"]["social_propagation_report"] = str(social_path)
    if agent_taxonomy_markdown:
        manifest["files"]["agent_taxonomy"] = str(agent_taxonomy_path)
    if policy_causal_chain is not None:
        manifest["files"]["policy_causal_chain"] = str(policy_causal_chain_path)
    if realism_metrics_csv:
        manifest["files"]["realism_metrics"] = str(realism_metrics_path)
    if competition_snapshot is not None:
        manifest["files"]["competition_mode_snapshot"] = str(competition_snapshot_path)
    if compliance_artifacts:
        for key, value in dict(compliance_artifacts.get("files", {})).items():
            manifest["files"][key] = str(value)
        if "manifest" in compliance_artifacts:
            manifest["competition_compliance"] = serializable_payload(dict(compliance_artifacts["manifest"]))

    manifest_path.write_text(json.dumps(manifest, ensure_ascii=False, indent=2, sort_keys=True), encoding="utf-8")
    return {
        "bundle_root": bundle_root,
        "manifest": manifest,
        "manifest_path": manifest_path,
        "realism_bundle": realism_bundle,
        "social_path": social_path if social_propagation_artifacts else None,
    }
